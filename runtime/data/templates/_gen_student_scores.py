# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "成绩"
headers = ["姓名", "班级", "语文", "数学", "英语", "科学"]
rows = [
    ["青禾", "东湖班", 88, 92, 85, 90],
    ["白鹭", "东湖班", 76, 68, 80, 72],
    ["赤松", "西岭班", 91, 95, 88, 93],
    ["暮云", "西岭班", 55, 60, 58, 62],
    ["晨露", "东湖班", 82, 79, 84, 77],
    ["南风", "西岭班", 70, 73, 69, 71],
]
thin = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
for col, h in enumerate(headers, 1):
    cell = ws.cell(1, col, h)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin
for r, row in enumerate(rows, 2):
    for c, val in enumerate(row, 1):
        cell = ws.cell(r, c, val)
        cell.border = thin
        if c >= 3:
            cell.alignment = Alignment(horizontal="center")
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 8
for col in ("C", "D", "E", "F"):
    ws.column_dimensions[col].width = 10

ws2 = wb.create_sheet("核算结果")
ws2["A1"] = "学生成绩核算"
ws2["A1"].font = Font(bold=True, size=14)
for col, h in enumerate(["姓名", "班级", "总分", "平均", "是否及格"], 1):
    cell = ws2.cell(2, col, h)
    cell.font = Font(bold=True)
    cell.border = thin
for col in ("A", "B", "C", "D", "E"):
    ws2.column_dimensions[col].width = 12

out = Path(__file__).with_name("学生各科分数.xlsx")
wb.save(out)
print("wrote", out, "bytes", out.stat().st_size)
